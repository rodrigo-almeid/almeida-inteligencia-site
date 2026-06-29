from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import extract, or_, and_
import io
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_RIGHT, TA_CENTER
from datetime import date

from backend.core import models
from backend.core.database import get_db
from backend.core.security import get_current_user, require_perfil

router = APIRouter(
    prefix="/relatorios",
    tags=["Relatórios e Exportação"],
    dependencies=[Depends(require_perfil("automacao_financeira"))]
)

MESES = ['Janeiro','Fevereiro','Março','Abril','Maio','Junho',
         'Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']

# ── Brand colors ──────────────────────────────────────────
C_BLACK      = colors.HexColor("#0A0F0A")
C_BG_MID     = colors.HexColor("#1A241C")
C_GREEN      = colors.HexColor("#3DDB82")
C_GREEN_DK   = colors.HexColor("#18A558")
C_GRAY       = colors.HexColor("#5C6B5E")
C_GRAY_L     = colors.HexColor("#A8B8AA")
C_RED        = colors.HexColor("#E24B4A")
C_YELLOW     = colors.HexColor("#F5A623")
C_WHITE      = colors.white
C_WHITE_COLD = colors.HexColor("#F7FAF7")
C_BORDER     = colors.HexColor("#D4E2D6")
C_ROW_ALT    = colors.HexColor("#F0F4F1")


def fmt_brl(v: float) -> str:
    s = f"{abs(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"R$ {s}" if v >= 0 else f"- R$ {s}"


def get_contas_filtradas(db, user_id, mes, ano):
    query = db.query(models.Conta).filter(models.Conta.user_id == user_id)
    if mes and ano:
        comp = f"{ano}-{mes:02d}"
        query = query.filter(or_(
            and_(models.Conta.competencia != None, models.Conta.competencia == comp),
            and_(models.Conta.competencia == None,
                 extract('year', models.Conta.vencimento) == ano,
                 extract('month', models.Conta.vencimento) == mes)
        ))
    return query.order_by(models.Conta.natureza, models.Conta.vencimento).all()


# ══════════════════════════════════════════════════════════
# EXCEL
# ══════════════════════════════════════════════════════════
@router.get("/excel/contas")
def exportar_excel_contas(
    mes: int = Query(None), ano: int = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    contas = get_contas_filtradas(db, current_user.id, mes, ano)
    cats = {c.id: c.nome for c in db.query(models.Categoria)
            .filter(models.Categoria.user_id == current_user.id).all()}

    titulo = f"{MESES[mes-1]} {ano}" if mes and ano else "Todas as contas"

    output = io.BytesIO()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contas"

    # ── Cabeçalho da marca ─────────────────────────────────
    ws.merge_cells("A1:G1")
    ws["A1"] = f"Almeida Inteligência — {titulo}"
    ws["A1"].font = Font(bold=True, size=14, color="3DDB82", name="Calibri")
    ws["A1"].fill = PatternFill("solid", fgColor="0A0F0A")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 34

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Exportado em {date.today().strftime('%d/%m/%Y')}  ·  {current_user.email}"
    ws["A2"].font = Font(size=9, color="A8B8AA", name="Calibri")
    ws["A2"].fill = PatternFill("solid", fgColor="0A0F0A")
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 6

    # ── Header da tabela ───────────────────────────────────
    headers = ["Descrição", "Vencimento", "Competência", "Categoria", "Natureza", "Status", "Valor (R$)"]
    green_bottom = Border(bottom=Side(style="medium", color="3DDB82"))
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=4, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=10)
        c.fill = PatternFill("solid", fgColor="1A241C")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = green_bottom
    ws.row_dimensions[4].height = 22

    receitas_total = 0.0
    despesas_total = 0.0
    thin_border = Border(bottom=Side(style="thin", color="D4E2D6"))

    for i, c in enumerate(contas, 5):
        alt_fill = PatternFill("solid", fgColor="F0F4F1") if i % 2 == 0 else None
        row_data = [
            c.descricao,
            c.vencimento.strftime("%d/%m/%Y") if c.vencimento else "",
            c.competencia or "",
            cats.get(c.categoria_id, "") if c.categoria_id else "",
            c.natureza or "",
            c.status or "",
            float(c.valor),
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border
            if alt_fill:
                cell.fill = alt_fill
            if col == 7:
                cell.number_format = '"R$" #,##0.00'
                cell.alignment = Alignment(horizontal="right")
            else:
                cell.alignment = Alignment(horizontal="left")
            if col == 5:
                cell.font = Font(name="Calibri", size=10,
                                 color="3DDB82" if val == "receita" else "E24B4A")
            if col == 6:
                clr = {"paga": "3DDB82", "pendente": "F5A623", "vencida": "E24B4A"}.get(val, "A8B8AA")
                cell.font = Font(name="Calibri", size=10, color=clr)

        if c.natureza == "receita":
            receitas_total += float(c.valor)
        else:
            despesas_total += float(c.valor)

    # ── Totais ─────────────────────────────────────────────
    tr = 4 + len(contas) + 2
    saldo = receitas_total - despesas_total
    saldo_cor = "18A558" if saldo >= 0 else "E24B4A"

    for offset, label, val, cor in [
        (0, "TOTAL RECEITAS", receitas_total, "18A558"),
        (1, "TOTAL DESPESAS", despesas_total, "E24B4A"),
        (2, "SALDO",           saldo,          saldo_cor),
    ]:
        top_border = Border(top=Side(style="medium" if offset==0 else "thin",
                                     color="D4E2D6"))
        lbl = ws.cell(row=tr+offset, column=6, value=label)
        lbl.font = Font(bold=True, size=10, color=cor, name="Calibri")
        lbl.border = top_border
        lbl.alignment = Alignment(horizontal="right")
        vl = ws.cell(row=tr+offset, column=7, value=val)
        vl.font = Font(bold=True, size=10 if offset < 2 else 12, color=cor, name="Calibri")
        vl.number_format = '"R$" #,##0.00'
        vl.border = top_border
        vl.alignment = Alignment(horizontal="right")

    # ── Larguras ───────────────────────────────────────────
    for col, w in enumerate([42, 14, 14, 22, 12, 12, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    # ── Rodapé ─────────────────────────────────────────────
    ws.sheet_view.showGridLines = False

    wb.save(output)
    output.seek(0)
    fname = f"contas_{ano}_{mes:02d}.xlsx" if mes and ano else "contas.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )


# ══════════════════════════════════════════════════════════
# PDF
# ══════════════════════════════════════════════════════════
@router.get("/pdf/contas")
def exportar_pdf_contas(
    mes: int = Query(None), ano: int = Query(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user)
):
    contas = get_contas_filtradas(db, current_user.id, mes, ano)
    cats = {c.id: c.nome for c in db.query(models.Categoria)
            .filter(models.Categoria.user_id == current_user.id).all()}

    titulo = f"Contas — {MESES[mes-1]} {ano}" if mes and ano else "Relatório de Contas"
    receitas_total = sum(float(c.valor) for c in contas if c.natureza == "receita")
    despesas_total = sum(float(c.valor) for c in contas if c.natureza == "despesa")
    saldo = receitas_total - despesas_total

    output = io.BytesIO()
    W, H = A4

    # ── Decorações de página ───────────────────────────────
    def draw_page(canvas, doc):
        canvas.saveState()

        # Barra superior escura
        canvas.setFillColor(C_BLACK)
        canvas.rect(0, H - 48*mm, W, 48*mm, fill=1, stroke=0)

        # Linha verde de acento
        canvas.setFillColor(C_GREEN)
        canvas.rect(0, H - 50*mm, W, 2*mm, fill=1, stroke=0)

        # Nome da marca
        canvas.setFont("Helvetica-Bold", 15)
        canvas.setFillColor(C_WHITE)
        canvas.drawString(20*mm, H - 18*mm, "Almeida")
        off = canvas.stringWidth("Almeida ", "Helvetica-Bold", 15)
        canvas.setFillColor(C_GREEN)
        canvas.drawString(20*mm + off, H - 18*mm, "Inteligência")

        # Subtítulo da marca
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(C_GRAY)
        canvas.drawString(20*mm, H - 25*mm, "Processos · IA · Resultados")

        # Título do relatório
        canvas.setFont("Helvetica-Bold", 11)
        canvas.setFillColor(C_WHITE)
        canvas.drawRightString(W - 20*mm, H - 19*mm, titulo)

        # Data de geração
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(C_GRAY_L)
        canvas.drawRightString(W - 20*mm, H - 26*mm,
                               f"Gerado em {date.today().strftime('%d/%m/%Y')}")

        # Rodapé
        canvas.setFillColor(C_BORDER)
        canvas.rect(20*mm, 13*mm, W - 40*mm, 0.4, fill=1, stroke=0)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(C_GRAY)
        canvas.drawString(20*mm, 8*mm, current_user.email)
        canvas.drawRightString(W - 20*mm, 8*mm, f"Página {doc.page}")

        canvas.restoreState()

    doc = SimpleDocTemplate(
        output, pagesize=A4,
        topMargin=56*mm, bottomMargin=22*mm,
        leftMargin=20*mm, rightMargin=20*mm,
    )

    elementos = []

    # ── Cards KPI ─────────────────────────────────────────
    kpi_style = TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), C_WHITE_COLD),
        ('BOX',        (0, 0), (-1, -1), 0.5, C_BORDER),
        ('INNERGRID',  (0, 0), (-1, -1), 0.5, C_BORDER),
        ('ALIGN',      (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN',     (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 10),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
        ('ROUNDEDCORNERS', [6]),
    ])

    def kpi_cell(label, value, cor):
        return Paragraph(
            f'<font color="#A8B8AA" size="7">{label}</font><br/>'
            f'<font color="{cor}" size="11"><b>{value}</b></font>',
            ParagraphStyle("kpi", fontName="Helvetica", alignment=TA_CENTER)
        )

    saldo_cor_hex = "#18A558" if saldo >= 0 else "#E24B4A"
    kpi_data = [[
        kpi_cell("RECEITAS",  fmt_brl(receitas_total), "#18A558"),
        kpi_cell("DESPESAS",  fmt_brl(despesas_total), "#E24B4A"),
        kpi_cell("SALDO",     fmt_brl(saldo),          saldo_cor_hex),
        kpi_cell("LANÇAMENTOS", str(len(contas)),       "#5C6B5E"),
    ]]
    usable = W - 40*mm
    kpi_table = Table(kpi_data, colWidths=[usable/4]*4)
    kpi_table.setStyle(kpi_style)
    elementos.append(kpi_table)
    elementos.append(Spacer(1, 6*mm))

    # ── Tabela de contas ───────────────────────────────────
    st_hdr = ParagraphStyle("th", fontName="Helvetica-Bold", fontSize=8.5,
                            textColor=C_GREEN, alignment=TA_CENTER)
    st_cel = ParagraphStyle("td", fontName="Helvetica", fontSize=8,
                            textColor=C_BLACK, alignment=TA_LEFT)
    st_val = ParagraphStyle("tv", fontName="Helvetica", fontSize=8,
                            textColor=C_BLACK, alignment=TA_RIGHT)
    st_val_b = ParagraphStyle("tvb", fontName="Helvetica-Bold", fontSize=9,
                               textColor=C_BLACK, alignment=TA_RIGHT)

    col_w = [54*mm, 18*mm, 18*mm, 22*mm, 16*mm, 15*mm, 27*mm]  # = 170mm

    rows = [[
        Paragraph(h, st_hdr) for h in
        ["Descrição", "Vencimento", "Competência", "Categoria", "Natureza", "Status", "Valor"]
    ]]

    for c in contas:
        nat_cor = "#3DDB82" if c.natureza == "receita" else "#E24B4A"
        sta_cor = {"paga":"#3DDB82","pendente":"#F5A623","vencida":"#E24B4A"}.get(c.status or "", "#A8B8AA")
        rows.append([
            Paragraph(esc_pdf(c.descricao), st_cel),
            Paragraph(c.vencimento.strftime("%d/%m/%Y") if c.vencimento else "—", st_cel),
            Paragraph(c.competencia or "—", st_cel),
            Paragraph(esc_pdf(cats.get(c.categoria_id,"") if c.categoria_id else "—"), st_cel),
            Paragraph(f'<font color="{nat_cor}">{c.natureza or "—"}</font>',
                      ParagraphStyle("nat", fontName="Helvetica", fontSize=8, alignment=TA_CENTER)),
            Paragraph(f'<font color="{sta_cor}">{c.status or "—"}</font>',
                      ParagraphStyle("sta", fontName="Helvetica", fontSize=8, alignment=TA_CENTER)),
            Paragraph(fmt_brl(float(c.valor)), st_val),
        ])

    # Linha de separação + totais
    empty = [Paragraph("", st_cel)] * 7
    rows.append(empty)
    for label, val, cor in [
        ("Receitas", receitas_total, "#18A558"),
        ("Despesas", despesas_total, "#E24B4A"),
        ("Saldo",    saldo,          saldo_cor_hex),
    ]:
        rows.append([
            Paragraph("", st_cel), Paragraph("", st_cel),
            Paragraph("", st_cel), Paragraph("", st_cel),
            Paragraph("", st_cel),
            Paragraph(f'<font color="{cor}"><b>{label}</b></font>',
                      ParagraphStyle("tl", fontName="Helvetica-Bold", fontSize=9, alignment=TA_RIGHT)),
            Paragraph(f'<font color="{cor}"><b>{fmt_brl(val)}</b></font>', st_val_b),
        ])

    tabela = Table(rows, colWidths=col_w, repeatRows=1)

    n = len(contas)
    ts = TableStyle([
        # Header
        ('BACKGROUND',    (0, 0), (-1, 0), C_BG_MID),
        ('LINEBELOW',     (0, 0), (-1, 0), 1.2, C_GREEN),
        ('TOPPADDING',    (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 9),
        # Linhas de dados
        ('FONTSIZE',      (0, 1), (-1, n), 8),
        ('TOPPADDING',    (0, 1), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 5),
        ('LINEBELOW',     (0, 1), (-1, n), 0.3, C_BORDER),
        # Separador de totais
        ('LINEABOVE',     (0, n+1), (-1, n+1), 1, C_BORDER),
        ('LINEABOVE',     (0, n+2), (-1, n+4), 0.3, C_BORDER),
        ('VALIGN',        (0, 0), (-1, -1), 'MIDDLE'),
    ])
    # Zebra
    for i in range(1, n+1):
        if i % 2 == 0:
            ts.add('BACKGROUND', (0, i), (-1, i), C_WHITE_COLD)
    tabela.setStyle(ts)
    elementos.append(tabela)

    doc.build(elementos, onFirstPage=draw_page, onLaterPages=draw_page)
    output.seek(0)
    fname = f"contas_{ano}_{mes:02d}.pdf" if mes and ano else "contas.pdf"
    return StreamingResponse(
        output, media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={fname}"}
    )


def esc_pdf(s: str) -> str:
    return str(s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
